#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""端到端回归测试：现场生成 fixture，断言转换器的关键不变量。

直接运行即可（无需 pytest）：
    python3 tests/test_convert.py

覆盖：
  1. 跨 run 拆词 —— 段落级聚合必须修复 头发→頭髮（逐节点会得到 頭發）
  2. Zip Slip   —— 含路径穿越条目的 .pptx 必须被拒绝且不写出文件
  3. 处理范围   —— 版式/母版/备注里的文字必须被转换
  4. 保真度     —— 页数/run 数/字体分布/媒体字节/压缩方式/条目顺序不变
  5. 幂等性     —— 二次转换不应再产生变化
  6. 文本文件   —— .md 转换与 .docx 拒绝
"""
import os
import re
import sys
import shutil
import zipfile
import tempfile
import subprocess
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS_DIR = os.path.join(ROOT, "skills", "opencc-s2hk", "scripts")
SCRIPT = os.path.join(SCRIPTS_DIR, "s2hk_convert.py")
PY = sys.executable
sys.path.insert(0, SCRIPTS_DIR)


def font_pairs(path):
    """返回文件中出现过的 (latin, ea) 字体组合集合。"""
    from pptx_core import font_profile
    return set(font_profile(path).keys())

passed = 0
failed = []


def check(cond, label, detail=""):
    global passed
    if cond:
        passed += 1
        print("  [pass] %s" % label)
    else:
        failed.append("%s %s" % (label, detail))
        print("  [FAIL] %s %s" % (label, detail))


def run(args):
    return subprocess.run([PY] + args, capture_output=True, text=True)


# ---------------------------------------------------------------- fixtures
def make_tricky(path):
    """跨 run 拆词 + 表格 + 版式/备注文字。"""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    p = tb.text_frame.paragraphs[0]
    # 「头发」「理发」被人为拆进不同 run（设计师加粗半个词的典型后果）
    for seg, bold in [("头", True), ("发", False), ("护理与", False),
                      ("理", True), ("发", False), ("服务", False)]:
        r = p.add_run()
        r.text = seg
        r.font.bold = bold
        r.font.size = Pt(24)
    p2 = tb.text_frame.add_paragraph()
    r = p2.add_run()
    r.text = "只身前往，干面还是汤面"
    t = s.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    t.cell(0, 0).text = "头发"
    t.cell(0, 1).text = "软件开发"
    t.cell(1, 0).text = "表里不一"
    t.cell(1, 1).text = "数据隔离"
    s.notes_slide.notes_text_frame.text = "备注页：项目总监统筹"
    prs.save(path)


def make_evil(src, dst):
    shutil.copy(src, dst)
    with zipfile.ZipFile(dst, "a") as z:
        z.writestr("../../../../tmp/S2HK_PWNED.txt", "zip slip")


def slide_texts(path):
    from pptx import Presentation
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:
                walk(sh.shapes)
            if sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    if p.runs:
                        out.append("".join(r.text for r in p.runs))
            if getattr(sh, "has_table", False):
                for row in sh.table.rows:
                    for c in row.cells:
                        out.append(c.text)
    for s in Presentation(path).slides:
        walk(s.shapes)
    return out


def main():
    try:
        from pptx import Presentation  # noqa: F401
        from opencc import OpenCC
    except ImportError as e:
        print("[skip] 缺少依赖 %s，跳过端到端测试" % getattr(e, "name", e))
        return 0

    tmp = tempfile.mkdtemp(prefix="s2hk-test-")
    try:
        src = os.path.join(tmp, "in.pptx")
        out = os.path.join(tmp, "out.pptx")
        make_tricky(src)

        print("\n[1] 跨 run 拆词（段落级聚合）")
        r = run([SCRIPT, src, "-o", out])
        check(r.returncode == 0, "转换成功", r.stderr[:200])
        texts = slide_texts(out)
        joined = "".join(texts)
        check("頭髮護理與理髮服務" in joined, "头发/理发 正确转为 頭髮/理髮",
              "实际: %s" % texts[:1])
        check("頭發" not in joined and "理發" not in joined,
              "没有出现逐节点转换的错误形式 頭發/理發")
        check("隻身" in joined, "只身 → 隻身（单 run 段落）")

        print("\n[2] Zip Slip 防护")
        evil = os.path.join(tmp, "evil.pptx")
        make_evil(src, evil)
        pwned = "/tmp/S2HK_PWNED.txt"
        if os.path.exists(pwned):
            os.remove(pwned)
        r = run([SCRIPT, evil, "-o", os.path.join(tmp, "evil_out.pptx")])
        check(r.returncode == 2, "恶意 pptx 被拒绝（exit 2）", "exit=%d" % r.returncode)
        check("路径穿越" in r.stderr or "绝对路径" in r.stderr,
              "给出可读的拒绝原因", r.stderr[:120])
        check(not os.path.exists(pwned), "没有文件逃逸到目标目录之外")
        r = run([SCRIPT, evil, "-o", os.path.join(tmp, "evil2.pptx"),
                 "--ea", "楷体"])
        check(r.returncode == 2 and "Traceback" not in r.stderr,
              "带字体参数时同样拒绝且无裸 traceback", "exit=%d" % r.returncode)

        print("\n[3] 处理范围（版式/母版/备注）")
        r = run([SCRIPT, src, "-o", out])
        cc = OpenCC("s2hk")
        A = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        from lxml import etree

        def residue(path, prefix=None):
            z = zipfile.ZipFile(path)
            n = 0
            for name in z.namelist():
                if not name.endswith(".xml"):
                    continue
                if prefix and not name.startswith(prefix):
                    continue
                try:
                    root = etree.fromstring(z.read(name))
                except Exception:
                    continue
                for t in root.iter(A):
                    if t.text and cc.convert(t.text) != t.text:
                        n += 1
            return n

        check(residue(out) == 0, "全文无残留简体（含版式/母版/备注）",
              "残留 %d 处" % residue(out))
        check(residue(src) > 0, "fixture 本身确实含待转文字（测试有效）")
        slides_only = os.path.join(tmp, "slides_only.pptx")
        run([SCRIPT, src, "-o", slides_only, "--scope", "slides"])
        check(residue(slides_only) > 0,
              "--scope slides 仅转正文页（旧行为仍可用）")
        check(residue(slides_only, "ppt/slides/") == 0,
              "--scope slides 下正文页本身已转干净")

        print("\n[4] 保真度")
        zi, zo = zipfile.ZipFile(src), zipfile.ZipFile(out)
        check(zi.namelist() == zo.namelist(), "zip 条目顺序与数量不变")
        ci = Counter(i.compress_type for i in zi.infolist())
        co = Counter(i.compress_type for i in zo.infolist())
        check(ci == co, "压缩方式分布不变（STORED 条目未被重新压缩）",
              "%s vs %s" % (dict(ci), dict(co)))
        from pptx import Presentation as P
        check(len(P(src).slides) == len(P(out).slides), "页数不变")
        check(len(slide_texts(src)) == len(slide_texts(out)), "文字块数不变")
        media = [n for n in zi.namelist() if n.startswith("ppt/media/")]
        check(all(zi.read(n) == zo.read(n) for n in media),
              "媒体文件字节级一致（%d 个）" % len(media))

        print("\n[5] 幂等性")
        out2 = os.path.join(tmp, "out2.pptx")
        run([SCRIPT, out, "-o", out2])
        check(slide_texts(out) == slide_texts(out2), "二次转换不再改变文字")

        print("\n[6] 文本文件与不支持的格式")
        md = os.path.join(tmp, "a.md")
        with open(md, "w", encoding="utf-8") as f:
            f.write("# 用户启用多租户\n\n数据隔离机制\n")
        r = run([SCRIPT, md])
        gen = os.path.join(tmp, "a.s2hk.md")
        check(r.returncode == 0 and os.path.isfile(gen), "Markdown 转换成功")
        if os.path.isfile(gen):
            body = open(gen, encoding="utf-8").read()
            check("數據隔離機制" in body, "Markdown 内容已转繁")
        docx = os.path.join(tmp, "b.docx")
        open(docx, "wb").write(b"PK\x03\x04fake")
        r = run([SCRIPT, docx])
        check(r.returncode == 2, "明确拒绝 .docx（exit 2）", "exit=%d" % r.returncode)

        print("\n[7] 内联文本")
        r = run([SCRIPT, "--text", "用户启用多租户数据隔离机制"])
        check(r.stdout.strip() == "用户啓用多租户數據隔離機制",
              "内联转换结果正确", r.stdout.strip())

        print("\n[8] 字体能力（已从 pptx-s2t-font-convert 并入）")
        fout = os.path.join(tmp, "font.pptx")
        r = run([SCRIPT, src, "-o", fout, "--ea", "楷体", "--latin", "Arial"])
        check(r.returncode == 0, "转换+改字体成功", r.stderr[:200])
        check("font props set" in r.stdout, "输出包含字体统计")
        pairs = font_pairs(fout)
        check(("Arial", "楷体") in pairs, "字体已写入 latin=Arial / ea=楷体",
              "实际: %s" % list(pairs)[:3])
        check(residue(fout) == 0, "改字体的同时文字仍完整转换")

        only = os.path.join(tmp, "fontonly.pptx")
        r = run([SCRIPT, src, "-o", only, "--mode", "none", "--ea", "楷体"])
        check(r.returncode == 0, "--mode none 仅改字体")
        check(slide_texts(only) == slide_texts(src), "--mode none 未改动任何文字")
        r = run([SCRIPT, src, "-o", os.path.join(tmp, "x.pptx"), "--mode", "none"])
        check(r.returncode == 2, "--mode none 缺字体参数时报错", "exit=%d" % r.returncode)

        print("\n[9] 内置质检 --verify")
        r = run([SCRIPT, src, "-o", os.path.join(tmp, "v.pptx"), "--verify"])
        check(r.returncode == 0 and "result            : PASS" in r.stdout,
              "纯转换质检 PASS", r.stdout[-200:])
        check("residue (full doc): 0" in r.stdout, "质检报告显示全文零残留")
        r = run([SCRIPT, src, "-o", os.path.join(tmp, "vf.pptx"),
                 "--verify", "--ea", "楷体"])
        check("font profile same : False" in r.stdout,
              "改字体时如实报告字体分布已变")
        check(r.returncode == 0 and "result            : PASS" in r.stdout,
              "改字体不应被判为质检失败（v1 的常见误导）")

        print("\n[10] --dry-run 与 --json")
        d = os.path.join(tmp, "dry.pptx")
        r = run([SCRIPT, src, "-o", d, "--dry-run"])
        check(r.returncode == 0 and not os.path.exists(d),
              "--dry-run 不写出文件")
        r = run([SCRIPT, src, "-o", os.path.join(tmp, "j.pptx"),
                 "--verify", "--json"])
        import json as _json
        try:
            data = _json.loads(r.stdout)
            check(data["verify"]["residue_count"] == 0 and data["verify_ok"],
                  "--json 输出可解析且质检通过")
            check(data["mode"] == "s2hk" and "slides" in data["scope"],
                  "--json 含 mode 与 scope")
        except ValueError:
            check(False, "--json 输出是合法 JSON", r.stdout[:200])
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

    print("\n%s" % ("-" * 50))
    print("通过 %d 项" % passed)
    if failed:
        print("失败 %d 项：" % len(failed))
        for f in failed:
            print("  - %s" % f)
        return 1
    print("全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
